"""``.xlsx`` input parser — every worksheet as a table (AI-INP-002, AI-INP-006).

The same deterministic reading as :mod:`.csvparser`, the same ``options`` and the same
:data:`~factory_intelligence.parsing.tabular.csvparser.HEADER_CATEGORY_RULES`: a workbook
differs from a CSV in *where* a cell is, so only the location seam changes.  Each entry is
located at the real A1 coordinate of the cell that produced its statement —
``Location.cell(sheet, "B3")`` — because AI-INP-009 must be able to list the position a
conflicting fact came from and a sheet name plus a row number is not a position a reviewer
can open the file at.

openpyxl is imported lazily and its absence is an error, never an empty fragment:
AI-INP-002 lists Excel as a *mandatory* input format and
:class:`factory_intelligence.errors.ParserDependencyMissing` exists exactly for the case
where the library for a mandatory format is not installed (AI-INP-003a's reasoning: a
silent empty NRR is worse than a loud failure).
"""

from __future__ import annotations

import hashlib
import io
from collections.abc import Iterator, Mapping, Sequence
from typing import Any

from factory_foundation.api import DataClassification

from factory_intelligence.enums import NrrStatus, ParserType
from factory_intelligence.errors import ParserDependencyMissing
from factory_intelligence.parsing.base import Fragment, Location, ParseError
from factory_intelligence.parsing.columnmapping import (
    TabularCollector,
    analyse_table,
    apply_mapping_gate,
    effective_headers,
    mapping_of,
)

#: Registered in ``factory_intelligence.registries.AI_ERROR_CODES`` (00 doc 6.5): the bytes
#: are not a workbook at all, so no cell — and therefore no location — exists.
INPUT_UNREADABLE = "AI_PARSER_INPUT_INVALID"


def _openpyxl() -> tuple[Any, Any]:
    """Import openpyxl, or refuse with the registered dependency error (AI-INP-002)."""
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ParserDependencyMissing(
            "openpyxl is not installed, so the mandatory .xlsx input format cannot be "
            "parsed; the workbook was not skipped and no empty NRR was produced"
        ) from exc
    return load_workbook, get_column_letter


class XlsxParser:
    """``.xlsx`` → NRR fragment, one table per worksheet, with the AI-INP-006 gate."""

    parser_type: ParserType = ParserType.XLSX
    parser_version: str = "1.0.0"

    def parse(
        self,
        content: bytes,
        *,
        source_artifact_id: str,
        classification: DataClassification,
        options: Mapping[str, Any],
    ) -> Fragment:
        load_workbook, get_column_letter = _openpyxl()
        fragment = Fragment(
            parser_type=self.parser_type,
            parser_version=self.parser_version,
            source_artifact_id=source_artifact_id,
            classification=classification,
        )
        workbook = _load(load_workbook, content)
        fragment.notes.update(_record_notes(content, options))
        try:
            selected = options.get("sheet") if isinstance(options.get("sheet"), str) else None
            sheets = [
                worksheet
                for worksheet in workbook.worksheets
                if selected is None or worksheet.title == selected
            ]
            if not sheets:
                raise ParseError(
                    INPUT_UNREADABLE,
                    f"the workbook has no worksheet named '{selected}'"
                    if selected
                    else "the workbook has no worksheets at all",
                )
            if selected is not None:
                fragment.notes["selectedSheet"] = selected
            has_header = bool(options.get("has_header", True))
            keys: list[str] = []
            for worksheet in sheets:
                self._read_sheet(
                    worksheet,
                    fragment=fragment,
                    get_column_letter=get_column_letter,
                    # per worksheet: a mapping confirmed for one sheet is not a mapping
                    # for the next (AI-INP-006)
                    mapping=mapping_of(options, worksheet.title),
                    has_header=has_header,
                    source_artifact_id=source_artifact_id,
                    keys=keys,
                )
        finally:
            close = getattr(workbook, "close", None)
            if callable(close):  # pragma: no branch - read-only workbooks always have one
                close()
        return fragment

    def _read_sheet(
        self,
        worksheet: Any,
        *,
        fragment: Fragment,
        get_column_letter: Any,
        mapping: Mapping[str, Any],
        has_header: bool,
        source_artifact_id: str,
        keys: list[str],
    ) -> None:
        sheet = str(worksheet.title)
        rows = [(number, values) for number, values in _material_rows(worksheet)]
        fragment.notes.setdefault("sheets", {})[sheet] = len(rows)
        if not rows:
            return
        if has_header:
            if len(rows) < 2:
                #: a header with no data rows states no requirement; recording the sheet in
                #: notes is honest, inventing an entry from a column name is not
                fragment.notes.setdefault("sheetsWithoutRows", []).append(sheet)
                return
            header_number, header_values = rows[0]
            raw_headers = [str(value if value is not None else "") for value in header_values]
            headers = effective_headers(raw_headers, mapping)
            body = rows[1:]
        else:
            width = max(len(values) for _number, values in rows)
            headers = effective_headers(
                [str(index) for index in range(1, width + 1)], mapping
            )
            body = rows
        #: every coordinate is a real cell of this sheet: the column letter of the cell that
        #: supplied the label, and the row number as the file counts it
        def location_of(row_number: int, column: int) -> Location:
            letter = str(get_column_letter(column + 1))
            return Location.cell(sheet, f"{letter}{row_number}")

        plan = analyse_table(headers, [values for _number, values in body])
        if plan is None:
            apply_mapping_gate(
                fragment,
                sheet=sheet,
                headers=headers,
                plan=None,
                unmatched=True,
                affected_keys=(),
                per_sheet=True,
            )
            return

        before = len(keys)
        TabularCollector(
            fragment=fragment,
            sheet=sheet,
            plan=plan,
            status=NrrStatus.CANDIDATE if not mapping else NrrStatus.CONFIRMED,
            source_artifact_id=source_artifact_id,
            location_of=location_of,
            keys=keys,
        ).collect(body)

        if not mapping:
            #: AI-INP-006 asks one question per sheet: two worksheets of one workbook are
            #: two tables and their mappings are two decisions a human can make separately
            apply_mapping_gate(
                fragment,
                sheet=sheet,
                headers=headers,
                plan=plan,
                unmatched=False,
                affected_keys=keys[before:],
                per_sheet=True,
            )


def _material_rows(worksheet: Any) -> Iterator[tuple[int, Sequence[Any]]]:
    for number, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
        row = list(values)
        if any(_present(value) for value in row):
            yield number, row


def _present(value: object) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    return True


def _load(load_workbook: Any, content: bytes) -> Any:
    """Read the workbook; a file openpyxl cannot open is an error, not an empty sheet.

    ``data_only=True`` because a requirements workbook's arithmetic is not the requirement:
    a formula cell yields its cached value, and where a formula has no cached value the cell
    reads empty rather than as the text of the formula.
    """
    try:
        return load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except ParserDependencyMissing:  # pragma: no cover - raised by the import guard only
        raise
    except Exception as exc:
        raise ParseError(
            INPUT_UNREADABLE, f"the .xlsx input could not be opened: {exc}"
        ) from exc


def _record_notes(content: bytes, options: Mapping[str, Any]) -> dict[str, Any]:
    """The AI-INP-004 fields a parse record carries besides parserType/parserVersion."""
    notes: dict[str, Any] = {
        "inputHash": "sha256:" + hashlib.sha256(content).hexdigest(),
        "inputBytes": len(content),
        "usedModel": False,
        "dataOnly": True,
    }
    filename = options.get("filename")
    if isinstance(filename, str) and filename.strip():
        notes["filename"] = filename.strip()
    return notes


__all__ = ["XlsxParser", "INPUT_UNREADABLE"]
